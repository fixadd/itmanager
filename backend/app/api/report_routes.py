import csv
import io
from datetime import date, datetime
from decimal import Decimal
from flask import Blueprint, jsonify, request, Response, send_file
from sqlalchemy import func, or_

from ..extensions import db
from ..models import Inventory, License, LicenseName, LicenseModel, StockItem, StockMovement, MaintenanceRecord, PurchaseRequest, PurchaseRequestItem, Personnel, ScrapRecord, ProductType, AuditLog, User, Factory, Department
from .auth_routes import permission_required, login_required

reports_bp = Blueprint("reports", __name__)


def counts(query, column):
    return [
        {"label": str(label or "Bilinmiyor"), "count": int(count)}
        for label, count in query.with_entities(column, func.count())
        .group_by(column).order_by(func.count().desc()).all()
    ]


def inventory_type_counts():
    rows = (
        db.session.query(ProductType.name, func.count(Inventory.id))
        .outerjoin(Inventory, Inventory.product_type_id == ProductType.id)
        .group_by(ProductType.id, ProductType.name)
        .order_by(func.count(Inventory.id).desc(), ProductType.name.asc())
        .all()
    )
    return [{"label": str(name or "Bilinmiyor"), "count": int(count)} for name, count in rows if count]


@reports_bp.get("/reports/summary")
@permission_required("reports.view")
def report_summary():
    def status_counts(model, status_column):
        rows = db.session.query(status_column, func.count(model.id)).group_by(status_column).all()
        return {str(label or "Bilinmiyor"): int(count) for label, count in rows}

    def inventory_faulty_count():
        return db.session.query(func.count(Inventory.id)).filter(
            Inventory.status.in_(["faulty", "arizali", "broken"])
        ).scalar() or 0

    inventory_status = status_counts(Inventory, Inventory.status)
    license_status = status_counts(License, License.status)
    stock_status = status_counts(StockItem, StockItem.status)
    maintenance_status = status_counts(MaintenanceRecord, MaintenanceRecord.status)
    request_status = status_counts(PurchaseRequest, PurchaseRequest.status)

    return jsonify({
        "inventory": {
            "total": sum(inventory_status.values()),
            "active": inventory_status.get("active", 0),
            "faulty": sum(inventory_status.get(k, 0) for k in ["faulty", "arizali", "broken"]),
            "maintenance": sum(inventory_status.get(k, 0) for k in ["maintenance", "service", "bakim"]),
            "scrapped": inventory_status.get("scrapped", 0),
            "by_status": [{"label": k, "count": v} for k, v in sorted(inventory_status.items(), key=lambda x: x[1], reverse=True)],
            "by_type": inventory_type_counts(),
        },
        "licenses": {
            "total": sum(license_status.values()),
            "active": license_status.get("active", 0),
            "expiring": license_status.get("expiring", 0),
            "expired": license_status.get("expired", 0),
            "scrapped": license_status.get("scrapped", 0),
            "by_status": [{"label": k, "count": v} for k, v in sorted(license_status.items(), key=lambda x: x[1], reverse=True)],
        },
        "stock": {
            "items": sum(stock_status.values()),
            "total_quantity": float(db.session.query(func.coalesce(func.sum(StockItem.quantity), 0)).scalar() or 0),
            "available": stock_status.get("available", 0),
            "by_status": [{"label": k, "count": v} for k, v in sorted(stock_status.items(), key=lambda x: x[1], reverse=True)],
        },
        "maintenance": {
            "total": sum(maintenance_status.values()),
            "pending": maintenance_status.get("pending", 0),
            "in_progress": maintenance_status.get("in_progress", 0),
            "completed": maintenance_status.get("completed", 0),
            "total_cost": float(db.session.query(func.coalesce(func.sum(MaintenanceRecord.cost), 0)).scalar() or 0),
            "by_status": [{"label": k, "count": v} for k, v in sorted(maintenance_status.items(), key=lambda x: x[1], reverse=True)],
        },
        "requests": {
            "total": sum(request_status.values()),
            "pending": request_status.get("pending", 0),
            "approved": request_status.get("approved", 0),
            "ordered": request_status.get("ordered", 0),
            "completed": request_status.get("completed", 0),
            "rejected": request_status.get("rejected", 0),
            "by_status": [{"label": k, "count": v} for k, v in sorted(request_status.items(), key=lambda x: x[1], reverse=True)],
        },
        "people": {
            "total": Personnel.query.count(),
            "active": Personnel.query.filter_by(active=True).count(),
            "inactive": Personnel.query.filter_by(active=False).count(),
        },
        "scrap": {
            "total": ScrapRecord.query.count(),
            "inventory": ScrapRecord.query.filter_by(source_type="inventory").count(),
            "license": ScrapRecord.query.filter_by(source_type="license").count(),
            "stock": ScrapRecord.query.filter_by(source_type="stock").count(),
        },
    })


@reports_bp.get("/dashboard/summary")
@login_required
def dashboard_summary():
    type_rows = inventory_type_counts()[:5]
    status_rows = counts(Inventory.query, Inventory.status)
    monthly_rows = (
        db.session.query(func.date_trunc("month", AuditLog.created_at), func.count(AuditLog.id))
        .group_by(func.date_trunc("month", AuditLog.created_at))
        .order_by(func.date_trunc("month", AuditLog.created_at).desc())
        .limit(4).all()
    )
    recent = (
        db.session.query(AuditLog, User.username)
        .outerjoin(User, User.id == AuditLog.actor_user_id)
        .order_by(AuditLog.created_at.desc())
        .limit(6).all()
    )
    inventory_status = {str(label or "Bilinmiyor"): int(count) for label, count in db.session.query(Inventory.status, func.count(Inventory.id)).group_by(Inventory.status).all()}
    request_status = {str(label or "Bilinmiyor"): int(count) for label, count in db.session.query(PurchaseRequest.status, func.count(PurchaseRequest.id)).group_by(PurchaseRequest.status).all()}
    maintenance_status = {str(label or "Bilinmiyor"): int(count) for label, count in db.session.query(MaintenanceRecord.status, func.count(MaintenanceRecord.id)).group_by(MaintenanceRecord.status).all()}
    movement_totals = dict(db.session.query(StockMovement.movement_type, func.coalesce(func.sum(StockMovement.quantity), 0)).group_by(StockMovement.movement_type).all())

    return jsonify({
        "inventory": {
            "total": sum(inventory_status.values()),
            "active": inventory_status.get("active", 0),
            "faulty": sum(inventory_status.get(k, 0) for k in ["faulty", "arizali", "broken"]),
            "maintenance": sum(inventory_status.get(k, 0) for k in ["maintenance", "service", "bakim"]),
            "scrapped": inventory_status.get("scrapped", 0),
            "by_status": status_rows,
            "by_type": type_rows,
        },
        "requests": {
            "pending": request_status.get("pending", 0),
            "approved": request_status.get("approved", 0),
            "ordered": request_status.get("ordered", 0),
            "completed": request_status.get("completed", 0),
        },
        "maintenance": {
            "pending": maintenance_status.get("pending", 0),
            "in_progress": maintenance_status.get("in_progress", 0),
            "service": maintenance_status.get("service", 0),
            "completed": maintenance_status.get("completed", 0),
        },
        "stock": {
            "total_quantity": float(db.session.query(func.coalesce(func.sum(StockItem.quantity), 0)).scalar() or 0),
            "in": float(movement_totals.get("in", 0) or 0),
            "out": float(movement_totals.get("out", 0) or 0),
        },
        "monthly_activity": [
            {"month": month.isoformat() if month else None, "count": int(count)}
            for month, count in reversed(monthly_rows)
        ],
        "recent_activity": [
            {"action": row.action, "entity_type": row.entity_type, "entity_id": row.entity_id,
             "actor": username or "Sistem", "created_at": row.created_at.isoformat() if row.created_at else None}
            for row, username in recent
        ],
    })


def _export_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _export_cell(value):
    value = _export_value(value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


@reports_bp.get("/reports/export")
@permission_required("reports.view")
def report_export():
    export_format = (request.args.get("format") or "xlsx").lower()
    if export_format not in {"xlsx", "csv"}:
        return jsonify({"error": "Desteklenmeyen dışa aktarma biçimi"}), 400

    search = (request.args.get("search") or "").strip()
    status = (request.args.get("status") or "").strip()
    factory_id = request.args.get("factory_id", type=int)
    department_id = request.args.get("department_id", type=int)
    sheets = {}

    inventory_q = Inventory.query
    if search:
        term = f"%{search}%"
        inventory_q = inventory_q.join(Inventory.product_type).join(Inventory.brand).filter(
            or_(Inventory.inventory_no.ilike(term), Inventory.serial_no.ilike(term),
                   Inventory.computer_name.ilike(term), ProductType.name.ilike(term))
        )
    if status:
        inventory_q = inventory_q.filter(Inventory.status == status)
    if factory_id:
        inventory_q = inventory_q.filter(Inventory.factory_id == factory_id)
    if department_id:
        inventory_q = inventory_q.filter(Inventory.department_id == department_id)
    inventory_rows = []
    for x in inventory_q.order_by(Inventory.id.asc()).all():
        inventory_rows.append({
            "Envanter No": x.inventory_no, "Barkod": x.barcode, "Seri No": x.serial_no,
            "Bilgisayar Adı": x.computer_name, "Makine No": x.machine_no, "IFS No": x.ifs_no,
            "Donanım Tipi": x.product_type.name if x.product_type else "",
            "Marka": x.brand.name if x.brand else "", "Model": x.model.name if x.model else "",
            "Fabrika": x.factory.name if x.factory else "", "Departman": x.department.name if x.department else "",
            "Atanan Personel": x.personnel.name if x.personnel else "", "Durum": x.status,
            "Not": x.note, "Oluşturulma": x.created_at, "Güncellenme": x.updated_at,
        })
    sheets["Envanter"] = inventory_rows

    stock_q = StockItem.query
    if search:
        term = f"%{search}%"
        stock_q = stock_q.join(StockItem.product_type).join(StockItem.brand).filter(
            db.or_(ProductType.name.ilike(term), StockItem.barcode.ilike(term), StockItem.note.ilike(term))
        )
    if status:
        stock_q = stock_q.filter(StockItem.status == status)
    stock_rows = []
    for x in stock_q.order_by(StockItem.id.asc()).all():
        stock_rows.append({
            "Barkod": x.barcode, "Donanım Tipi": x.product_type.name if x.product_type else "",
            "Marka": x.brand.name if x.brand else "", "Model": x.model.name if x.model else "",
            "Miktar": x.quantity, "Birim": x.unit, "Durum": x.status, "Not": x.note,
            "Oluşturulma": x.created_at, "Güncellenme": x.updated_at,
        })
    sheets["Stok"] = stock_rows

    license_q = License.query
    if search:
        term = f"%{search}%"
        license_q = license_q.join(License.license_name).filter(
            db.or_(LicenseName.name.ilike(term), License.barcode.ilike(term),
                   License.email.ilike(term), License.note.ilike(term))
        )
    if status:
        license_q = license_q.filter(License.status == status)
    license_rows = []
    for x in license_q.order_by(License.id.asc()).all():
        license_rows.append({
            "Barkod": x.barcode, "Lisans Adı": x.license_name.name if x.license_name else "",
            "Lisans Modeli": x.license_model.name if x.license_model else "",
            "Lisans Türü": x.license_type, "Durum": x.status,
            "Atanan Personel": x.personnel.name if x.personnel else "",
            "Bağlı Envanter No": x.inventory.inventory_no if x.inventory else "",
            "E-posta": x.email, "Başlangıç": x.starts_at, "Bitiş": x.expires_at,
            "Not": x.note, "Oluşturulma": x.created_at, "Güncellenme": x.updated_at,
        })
    # License keys and passwords are intentionally excluded from exports.
    sheets["Lisanslar"] = license_rows

    movement_rows = []
    for x in StockMovement.query.order_by(StockMovement.created_at.asc(), StockMovement.id.asc()).all():
        movement_rows.append({
            "Tarih": x.created_at, "Stok Barkodu": x.stock_item.barcode if x.stock_item else "",
            "Hareket": x.movement_type, "Miktar": x.quantity, "Birim": x.unit,
            "Personel": x.personnel.name if x.personnel else "",
            "Envanter No": x.inventory.inventory_no if x.inventory else "", "Not": x.note,
        })
    sheets["Stok Hareketleri"] = movement_rows

    request_rows = []
    request_query = PurchaseRequest.query\n    if status:\n        request_query = request_query.filter(PurchaseRequest.status == status)\n    if search:\n        term = f"%{search}%"\n        request_query = request_query.outerjoin(PurchaseRequest.requester).filter(or_(PurchaseRequest.request_no.ilike(term), Personnel.name.ilike(term)))\n    for x in request_query.order_by(PurchaseRequest.requested_at.asc(), PurchaseRequest.id.asc()).all():
        creator_audits = AuditLog.query.filter(
            AuditLog.entity_type == "purchase_request", AuditLog.entity_id == x.id
        ).order_by(AuditLog.created_at.asc()).all()
        creator_names = []
        for audit in creator_audits:
            if audit.actor_user_id:
                actor = db.session.get(User, audit.actor_user_id)
                if actor:
                    creator_names.append(actor.full_name)
        for item in x.items:
            request_rows.append({
                "Talep No": x.request_no, "Talep Sahibi": x.requester.name if x.requester else "",
                "Talebi Açan Kullanıcı": creator_names[0] if creator_names else "",
                "Fabrika": x.factory.name if x.factory else "", "Departman": x.department.name if x.department else "",
                "Durum": x.status, "Talep Tarihi": x.requested_at, "Onay Tarihi": x.approved_at,
                "Tamamlanma Tarihi": x.completed_at, "Ürün Türü": item.product_type,
                "Donanım Tipi": item.device_type, "Marka": item.brand, "Model": item.model,
                "Talep Miktarı": item.quantity, "Teslim Alınan": item.received_quantity,
                "Birim": item.unit, "Açıklama": item.description, "Talep Notu": x.note,
            })
    sheets["Satın Alma Talepleri"] = request_rows

    history_rows = []
    audits = AuditLog.query.order_by(AuditLog.created_at.asc(), AuditLog.id.asc()).all()
    actor_ids = {a.actor_user_id for a in audits if a.actor_user_id}
    actors = {u.id: u.full_name for u in User.query.filter(User.id.in_(actor_ids)).all()} if actor_ids else {}
    for a in audits:
        details = a.details if isinstance(a.details, dict) else {}
        history_rows.append({
            "Tarih": a.created_at, "İşlemi Yapan": actors.get(a.actor_user_id, "Sistem"),
            "İşlem": a.action, "Kayıt Türü": a.entity_type, "Kayıt ID": a.entity_id,
            "Talep No": details.get("request_no") or details.get("request_number") or "",
            "Talep Sahibi": details.get("requester_name") or "",
            "Talebi Açan Kullanıcı": details.get("request_creator_name") or "",
            "Detay": str(details),
        })
    sheets["İşlem Geçmişi"] = history_rows

    if export_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(["Modül", "Kayıt Verisi (JSON)"])
        for sheet_name, rows in sheets.items():
            if not rows:
                writer.writerow([sheet_name, '{"bilgi":"Kayıt bulunamadı"}'])
                continue
            for row in rows:
                safe_row = {key: _export_cell(value) for key, value in row.items()}
                writer.writerow([sheet_name, __import__("json").dumps(safe_row, ensure_ascii=False, default=str)])
        data = "\ufeff" + output.getvalue()
        return Response(data, mimetype="text/csv; charset=utf-8",
                        headers={"Content-Disposition": 'attachment; filename="itmanager-rapor.csv"'})

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"error": "XLSX dışa aktarma bileşeni kurulu değil"}), 500
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        ws = workbook.create_sheet(sheet_name[:31])
        headers = list(rows[0].keys()) if rows else ["Kayıt"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24476B")
        if rows:
            for row in rows:
                ws.append([_export_cell(row.get(h)) for h in headers])
        else:
            ws.append(["Kayıt bulunamadı"])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx, header in enumerate(headers, 1):
            values = [str(header)] + [str(ws.cell(row=i, column=col_idx).value or "") for i in range(2, min(ws.max_row, 200) + 1)]
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max(map(len, values), default=10) + 2, 12), 42)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="itmanager-rapor.xlsx")
